import io
from pathlib import Path
from typing import List, Tuple
import pypdf
import docx

class DocumentParseError(Exception):
    pass

def parse_document(file_bytes: bytes, filename: str) -> List[Tuple[int, str]]:
    """
    Parses PDF, DOCX, or TXT file bytes into a list of (page_number, text_content).
    Handles empty/scanned detection.
    """
    ext = Path(filename).suffix.lower()
    pages_data: List[Tuple[int, str]] = []

    if ext == ".pdf":
        try:
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            if reader.is_encrypted:
                raise DocumentParseError("Encrypted PDFs are not supported.")
            for idx, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                pages_data.append((idx + 1, text.strip()))
        except Exception as e:
            if isinstance(e, DocumentParseError):
                raise
            raise DocumentParseError(f"Failed to parse PDF {filename}: {str(e)}")

    elif ext == ".docx":
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_str = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_str:
                        full_text.append(row_str)
            pages_data.append((1, "\n".join(full_text)))
        except Exception as e:
            raise DocumentParseError(f"Failed to parse DOCX {filename}: {str(e)}")

    elif ext in [".txt", ".md"]:
        try:
            text = file_bytes.decode("utf-8", errors="replace")
            pages_data.append((1, text.strip()))
        except Exception as e:
            raise DocumentParseError(f"Failed to parse TXT {filename}: {str(e)}")
    else:
        raise DocumentParseError(f"Unsupported file format: {ext}. Allowed: .pdf, .docx, .txt, .md")

    # Detect scanned or empty docs
    total_text = "".join(t for _, t in pages_data)
    if not total_text.strip():
        raise DocumentParseError(f"No extractable text found in {filename}. It may be scanned or empty.")

    return pages_data
